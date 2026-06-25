import sys
import subprocess

# Ensure openpyxl is installed
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl not found. Installing openpyxl...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

# Create workbook and sheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Pilot Registry"

# Enable grid lines visibility
ws.views.sheetView[0].showGridLines = True

# Define columns
headers = [
    "License Key",
    "Photographer Name",
    "Company",
    "Country",
    "Date Issued",
    "Date Activated",
    "Photos Processed",
    "Hours Saved",
    "Accuracy Rating (1-10)",
    "Would Pay? (Y/N)",
    "Feedback Status",
    "Bug Reports",
    "Renewal Candidate"
]

# Write headers
ws.append(headers)

# Populate rows for QC-PILOT-001 through QC-PILOT-020
for i in range(1, 21):
    ws.append([f"QC-PILOT-{i:03d}", "", "", "", "", "", "", "", "", "", "Unsent", "", ""])

# Apply Styling
font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
fill_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid") # Dark Blue
align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center")

thin_side = Side(border_style="thin", color="D9D9D9")
border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

font_data = Font(name="Segoe UI", size=10, color="000000")
font_key = Font(name="Consolas", size=10, bold=True, color="1F497D")

# Header styling
for col_idx in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_idx)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center
    cell.border = border_cell
ws.row_dimensions[1].height = 28

# Data row styling
for r_idx in range(2, 22):
    ws.row_dimensions[r_idx].height = 20
    # License Key column (Col 1)
    cell_key = ws.cell(row=r_idx, column=1)
    cell_key.font = font_key
    cell_key.alignment = align_center
    cell_key.border = border_cell
    
    # Other columns
    for col_idx in range(2, len(headers) + 1):
        cell = ws.cell(row=r_idx, column=col_idx)
        cell.font = font_data
        cell.border = border_cell
        # Center align short status/metric columns
        if col_idx in [5, 6, 7, 8, 9, 10, 11, 13]:
            cell.alignment = align_center
        else:
            cell.alignment = align_left

# Autofit column widths
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        if cell.value:
            max_len = max(max_len, len(str(cell.value)))
    # Add a bit of padding
    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

# Specific manual adjustment for wrapped headers
ws.column_dimensions['A'].width = 16 # License Key
ws.column_dimensions['B'].width = 22 # Name
ws.column_dimensions['C'].width = 22 # Company
ws.column_dimensions['G'].width = 18 # Photos Processed
ws.column_dimensions['H'].width = 15 # Hours Saved
ws.column_dimensions['I'].width = 22 # Accuracy Rating
ws.column_dimensions['K'].width = 18 # Feedback Status

# Save workbook
output_path = "PILOT_TRACKER.xlsx"
wb.save(output_path)
print(f"Successfully created styled {output_path}!")
