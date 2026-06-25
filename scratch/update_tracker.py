import openpyxl
import sqlite3
from datetime import datetime

# Connect to DB and fetch active keys
conn = sqlite3.connect("server_activations.db")
cursor = conn.cursor()
cursor.execute("SELECT license_key, name, company, country, activated_at FROM licenses WHERE status='active'")
active_records = {r[0]: r for r in cursor.fetchall()}
conn.close()

if not active_records:
    print("No active licenses found in database.")
    exit(0)

# Load Excel workbook
wb = openpyxl.load_workbook("PILOT_TRACKER.xlsx")
ws = wb.active

# Iterate through rows and update
updated = 0
for row_idx in range(2, ws.max_row + 1):
    key = ws.cell(row=row_idx, column=1).value
    if key in active_records:
        rec = active_records[key]
        ws.cell(row=row_idx, column=2, value=rec[1]) # Photographer Name
        ws.cell(row=row_idx, column=3, value=rec[2]) # Company
        ws.cell(row=row_idx, column=4, value=rec[3]) # Country
        
        # Format Dates
        act_date_str = rec[4].split("T")[0] if rec[4] else ""
        ws.cell(row=row_idx, column=5, value=act_date_str) # Date Issued (same as active for this run)
        ws.cell(row=row_idx, column=6, value=act_date_str) # Date Activated
        ws.cell(row=row_idx, column=11, value="Activated") # Feedback Status
        updated += 1
        print(f"Updated {key} in PILOT_TRACKER.xlsx")

if updated > 0:
    wb.save("PILOT_TRACKER.xlsx")
    print("Successfully saved PILOT_TRACKER.xlsx!")
else:
    print("No matches found to update.")
